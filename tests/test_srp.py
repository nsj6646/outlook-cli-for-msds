import os
import sys
import unittest
import openpyxl
import tempfile

# 프로젝트 루트를 path에 추가
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from core.distribution import ExcelParser, HtmlTableRenderer, TemplateEngine

class TestHtmlTableRenderer(unittest.TestCase):
    def test_renderer_renders_valid_html_table(self):
        """HtmlTableRenderer가 주어진 표 데이터로 미려한 인라인 CSS HTML 표를 렌더링하는지 확인"""
        headers = ["품목", "수량", "단가"]
        rows = [
            {"품목": "사과", "수량": 10, "단가": 1000},
            {"품목": "바나나", "수량": 5, "단가": 2000}
        ]
        
        renderer = HtmlTableRenderer(headers)
        html = renderer.render(rows)
        
        self.assertIn("<table", html)
        self.assertIn("border-collapse: collapse;", html)
        self.assertIn("사과", html)
        self.assertIn("1000", html)
        self.assertIn("바나나", html)
        self.assertIn("2000", html)

    def test_renderer_returns_empty_string_for_empty_rows(self):
        """HtmlTableRenderer에 행 데이터가 없을 때 빈 문자열을 반환하는지 확인"""
        headers = ["품목", "수량"]
        renderer = HtmlTableRenderer(headers)
        html = renderer.render([])
        self.assertEqual(html, "")

    def test_renderer_returns_empty_string_for_empty_headers(self):
        """HtmlTableRenderer에 헤더가 없을 때 빈 문자열을 반환하는지 확인"""
        rows = [{"품목": "사과"}]
        renderer = HtmlTableRenderer([])
        html = renderer.render(rows)
        self.assertEqual(html, "")


class TestTemplateEngine(unittest.TestCase):
    def test_template_engine_replaces_placeholders_case_insensitively(self):
        """TemplateEngine이 {{변수명}} 플레이스홀더를 대소문자 구분 없이 올바르게 치환하는지 확인"""
        template = "안녕하세요, {{Name}}님. 귀하의 이메일은 {{email}}입니다."
        replace_dict = {
            "name": "홍길동",
            "email": "gildong@example.com"
        }
        
        engine = TemplateEngine(template)
        rendered = engine.render(replace_dict)
        
        self.assertEqual(rendered, "안녕하세요, 홍길동님. 귀하의 이메일은 gildong@example.com입니다.")

    def test_template_engine_retains_unknown_placeholders(self):
        """replace_dict에 매칭되는 키가 없을 때 플레이스홀더 원본을 그대로 유지하는지 확인"""
        template = "안녕하세요, {{Name}}님. {{Unknown}} 값은 그대로 둡니다."
        replace_dict = {
            "name": "홍길동"
        }
        
        engine = TemplateEngine(template)
        rendered = engine.render(replace_dict)
        
        self.assertEqual(rendered, "안녕하세요, 홍길동님. {{Unknown}} 값은 그대로 둡니다.")

    def test_template_engine_returns_empty_string_for_empty_template(self):
        """템플릿 내용이 비어 있을 때 빈 문자열을 반환하는지 확인"""
        engine = TemplateEngine("")
        rendered = engine.render({"key": "val"})
        self.assertEqual(rendered, "")


class TestExcelParser(unittest.TestCase):
    def setUp(self):
        # 테스트용 임시 엑셀 파일 생성
        self.temp_dir = tempfile.TemporaryDirectory()
        self.excel_path = os.path.join(self.temp_dir.name, "test_mail_data.xlsx")
        
        wb = openpyxl.Workbook()
        # 1. Recipients 시트 작성
        ws_rec = wb.active
        ws_rec.title = "Recipients"
        ws_rec.append(["To", "Cc", "Subject", "Body", "Name", "Attachments"])
        ws_rec.append(["test1@example.com", "cc1@example.com", "제목1", "내용1", "홍길동", ""])
        ws_rec.append(["test2@example.com", "", "제목2", "", "이순신", "attach1.txt;attach2.txt"])
        
        # 2. TableData 시트 작성
        ws_tbl = wb.create_sheet(title="TableData")
        ws_tbl.append(["Email", "상품명", "금액"])
        ws_tbl.append(["test1@example.com", "노트북", 1500000])
        ws_tbl.append(["test1@example.com", "마우스", 50000])
        ws_tbl.append(["test2@example.com", "키보드", 120000])
        
        wb.save(self.excel_path)
        wb.close()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_parser_extracts_recipients_metadata_correctly(self):
        """ExcelParser가 Recipients 시트의 수신자 메타데이터와 변수 사전을 정확하게 로드하는지 확인"""
        parser = ExcelParser(self.excel_path)
        recipients = parser.get_recipients()
        
        self.assertEqual(len(recipients), 2)
        
        # 첫 번째 수신자 검증
        rec1 = recipients[0]
        self.assertEqual(rec1["to"], "test1@example.com")
        self.assertEqual(rec1["cc"], "cc1@example.com")
        self.assertEqual(rec1["subject"], "제목1")
        self.assertEqual(rec1["body"], "내용1")
        self.assertEqual(rec1["replace_dict"]["name"], "홍길동")
        self.assertEqual(rec1["attachments"], [])
        
        # 두 번째 수신자 검증
        rec2 = recipients[1]
        self.assertEqual(rec2["to"], "test2@example.com")
        self.assertEqual(rec2["cc"], "")
        self.assertEqual(rec2["subject"], "제목2")
        self.assertEqual(rec2["body"], "")
        self.assertEqual(rec2["replace_dict"]["name"], "이순신")
        
        # 첨부파일 경로 검증 (절대 경로 변환으로 끝부분 비교)
        self.assertTrue(len(rec2["attachments"]) == 2)
        self.assertTrue(rec2["attachments"][0].endswith("attach1.txt"))
        self.assertTrue(rec2["attachments"][1].endswith("attach2.txt"))

    def test_parser_groups_table_data_by_email_correctly(self):
        """ExcelParser가 TableData 시트의 데이터를 이메일 주소별(소문자)로 그룹화하여 캐싱하는지 확인"""
        parser = ExcelParser(self.excel_path)
        
        # test1@example.com의 표 데이터 검증
        tbl1 = parser.get_table_data("test1@example.com")
        self.assertEqual(len(tbl1), 2)
        self.assertEqual(tbl1[0]["상품명"], "노트북")
        self.assertEqual(tbl1[0]["금액"], 1500000)
        self.assertEqual(tbl1[1]["상품명"], "마우스")
        
        # 대소문자 무관 검색 확인
        tbl1_upper = parser.get_table_data("TEST1@EXAMPLE.COM")
        self.assertEqual(len(tbl1_upper), 2)
        
        # test2@example.com의 표 데이터 검증
        tbl2 = parser.get_table_data("test2@example.com")
        self.assertEqual(len(tbl2), 1)
        self.assertEqual(tbl2[0]["상품명"], "키보드")

    def test_parser_extracts_table_headers_excluding_email(self):
        """ExcelParser가 TableData의 헤더 목록 중 Email을 제외하고 정상 추출하는지 확인"""
        parser = ExcelParser(self.excel_path)
        headers = parser.get_table_headers()
        self.assertEqual(headers, ["상품명", "금액"])

    def test_parser_raises_value_error_for_missing_recipients_sheet(self):
        """Recipients 시트가 누락되었을 때 ValueError가 발생하는지 확인"""
        bad_excel = os.path.join(self.temp_dir.name, "bad_excel.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "NotRecipients"
        wb.save(bad_excel)
        wb.close()
        
        with self.assertRaises(ValueError) as ctx:
            ExcelParser(bad_excel)
        self.assertIn("Recipients", str(ctx.exception))

    def test_parser_raises_value_error_for_missing_to_column(self):
        """Recipients 시트에 필수 To 컬럼이 없을 때 ValueError가 발생하는지 확인"""
        bad_excel = os.path.join(self.temp_dir.name, "bad_excel2.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Recipients"
        ws.append(["Cc", "Subject", "Body"])  # To 누락
        wb.save(bad_excel)
        wb.close()
        
        with self.assertRaises(ValueError) as ctx:
            ExcelParser(bad_excel)
        self.assertIn("To", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
