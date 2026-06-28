import os
import re
import datetime
import logging
from dataclasses import dataclass, field
from typing import List, Optional
import openpyxl

@dataclass
class MailJob:
    to_addr: str
    subject: str
    body: str
    cc_addr: Optional[str] = None
    bcc_addr: Optional[str] = None
    attachments: List[str] = field(default_factory=list)
    from_addr: Optional[str] = None
    deferred_time: Optional[str] = None


class ExcelParser:
    """엑셀 파일을 읽고, Recipients 시트와 TableData 시트를 로딩하여 수신자 메타데이터와 표 데이터를 파싱 및 캐싱하는 책임."""
    def __init__(self, excel_path: str):
        self.excel_path = excel_path
        self.recipients_data = []
        self.table_groups = {}
        self.table_headers = []
        self.rec_headers = []
        self._parse()

    def _parse(self):
        if not os.path.exists(self.excel_path):
            raise FileNotFoundError(f"지정한 엑셀 파일을 찾을 수 없습니다: {os.path.abspath(self.excel_path)}")
            
        wb = None
        try:
            wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        except Exception as e:
            raise ValueError(f"엑셀 파일을 로드하는 중 오류가 발생했습니다: {str(e)}")
            
        try:
            recipients_sheet_name = None
            tabledata_sheet_name = None
            
            for name in wb.sheetnames:
                if name.lower() == "recipients":
                    recipients_sheet_name = name
                elif name.lower() == "tabledata":
                    tabledata_sheet_name = name
                    
            if not recipients_sheet_name:
                raise ValueError("엑셀 파일에 'Recipients' 시트가 존재하지 않습니다.")
                
            ws_recipients = wb[recipients_sheet_name]
            
            # TableData 시트가 있으면 표 데이터 캐싱
            if tabledata_sheet_name:
                ws_tabledata = wb[tabledata_sheet_name]
                raw_headers = [cell.value for cell in ws_tabledata[1]]
                
                email_col_idx = None
                for idx, h in enumerate(raw_headers, 1):
                    if h and str(h).strip().lower() == "email":
                        email_col_idx = idx
                        break
                        
                if email_col_idx is None:
                    raise ValueError("'TableData' 시트에 필수 'Email' 컬럼이 누락되었습니다.")
                    
                self.table_headers = [str(h).strip() for idx, h in enumerate(raw_headers, 1) if idx != email_col_idx and h]
                
                for r_idx in range(2, ws_tabledata.max_row + 1):
                    email_val = ws_tabledata.cell(row=r_idx, column=email_col_idx).value
                    if not email_val:
                        continue
                    email_key = str(email_val).strip().lower()
                    
                    row_data = {}
                    for idx, h in enumerate(raw_headers, 1):
                        if idx == email_col_idx or not h:
                            continue
                        row_data[str(h).strip()] = ws_tabledata.cell(row=r_idx, column=idx).value
                        
                    if email_key not in self.table_groups:
                        self.table_groups[email_key] = []
                    self.table_groups[email_key].append(row_data)

            self.rec_headers = [cell.value for cell in ws_recipients[1]]
            rec_mapping = {}
            
            for idx, cell_value in enumerate(self.rec_headers, 1):
                if cell_value:
                    h_name = str(cell_value).strip().lower()
                    rec_mapping[h_name] = idx
                    
            if "to" not in rec_mapping:
                raise ValueError("'Recipients' 시트에 필수 'To' 컬럼이 누락되었습니다.")
                
            max_row = ws_recipients.max_row
            for row_idx in range(2, max_row + 1):
                to_col = rec_mapping["to"]
                to_value = ws_recipients.cell(row=row_idx, column=to_col).value
                if not to_value:
                    continue
                    
                to_addr = str(to_value).strip()
                
                cc_addr = ""
                if "cc" in rec_mapping:
                    val = ws_recipients.cell(row=row_idx, column=rec_mapping["cc"]).value
                    cc_addr = str(val).strip() if val else ""
                    
                bcc_addr = ""
                if "bcc" in rec_mapping:
                    val = ws_recipients.cell(row=row_idx, column=rec_mapping["bcc"]).value
                    bcc_addr = str(val).strip() if val else ""
                    
                subject = ""
                if "subject" in rec_mapping:
                    val = ws_recipients.cell(row=row_idx, column=rec_mapping["subject"]).value
                    subject = str(val).strip() if val else ""
                    
                attachments = []
                if "attachments" in rec_mapping:
                    val = ws_recipients.cell(row=row_idx, column=rec_mapping["attachments"]).value
                    if val:
                        paths = str(val).split(";")
                        for p in paths:
                            p_clean = p.strip().strip('"').strip("'")
                            if p_clean:
                                attachments.append(os.path.abspath(p_clean))
                                
                body_val = ""
                if "body" in rec_mapping:
                    val = ws_recipients.cell(row=row_idx, column=rec_mapping["body"]).value
                    body_val = str(val) if val is not None else ""
                
                from_addr = ""
                from_col_idx = rec_mapping.get("from") or rec_mapping.get("sender")
                if from_col_idx:
                    val = ws_recipients.cell(row=row_idx, column=from_col_idx).value
                    from_addr = str(val).strip() if val else ""

                deferred_time_val = None
                deferred_col_idx = (
                    rec_mapping.get("deferredtime") or 
                    rec_mapping.get("deferred_time") or 
                    rec_mapping.get("deferred") or 
                    rec_mapping.get("예약시간") or 
                    rec_mapping.get("예약발송")
                )
                if deferred_col_idx:
                    val = ws_recipients.cell(row=row_idx, column=deferred_col_idx).value
                    if val is not None:
                        if isinstance(val, str):
                            deferred_time_val = val.strip()
                        else:
                            deferred_time_val = val
                
                row_replace_dict = {}
                for idx, col_name in enumerate(self.rec_headers, 1):
                    if not col_name:
                        continue
                    val = ws_recipients.cell(row=row_idx, column=idx).value
                    val_str = str(val) if val is not None else ""
                    row_replace_dict[str(col_name).strip().lower()] = val_str

                self.recipients_data.append({
                    "to": to_addr,
                    "cc": cc_addr,
                    "bcc": bcc_addr,
                    "subject": subject,
                    "attachments": attachments,
                    "body": body_val,
                    "from": from_addr,
                    "deferred_time": deferred_time_val,
                    "replace_dict": row_replace_dict
                })
        finally:
            if wb:
                wb.close()

    def get_recipients(self):
        return self.recipients_data

    def get_table_data(self, email: str):
        email_key = email.strip().lower()
        return self.table_groups.get(email_key, [])

    def get_table_headers(self):
        return self.table_headers


class HtmlTableRenderer:
    """특정 이메일에 귀속된 표 데이터를 바탕으로 인라인 CSS가 가미된 미려한 HTML 표 문자열을 빌드하는 책임."""
    def __init__(self, table_headers: list):
        self.table_headers = table_headers
        self.table_style = "border-collapse: collapse; border: none; font-family: '맑은 고딕', 'Malgun Gothic', sans-serif; font-size: 11.0pt; margin: 15px 0;"
        self.th_style = "border: solid windowtext 1.0pt; padding: 0cm 5.4pt 0cm 5.4pt; font-weight: bold; text-align: left; color: #000000; background-color: transparent;"
        self.td_style = "border: solid windowtext 1.0pt; padding: 0cm 5.4pt 0cm 5.4pt; color: #000000;"

    def render(self, rows: list) -> str:
        if not rows or not self.table_headers:
            return ""
        
        html_table = f'<table style="{self.table_style}">\n<thead>\n<tr>\n'
        for th in self.table_headers:
            html_table += f'  <th style="{self.th_style}">{th}</th>\n'
        html_table += '</tr>\n</thead>\n<tbody>\n'
        
        for row in rows:
            html_table += '<tr>\n'
            for th in self.table_headers:
                val = row.get(th, "")
                val_str = str(val) if val is not None else ""
                html_table += f'  <td style="{self.td_style}">{val_str}</td>\n'
            html_table += '</tr>\n'
        html_table += '</tbody>\n</table>'
        return html_table


class TemplateEngine:
    """이메일 템플릿 텍스트에 대해 대소문자 구분 없이 {{변수명}} 플레이스홀더를 치환하는 책임."""
    def __init__(self, template_content: str):
        self.template_content = template_content
        self.pattern = re.compile(r'\{\{\s*(.*?)\s*\}\}')

    def render(self, replace_dict: dict) -> str:
        if not self.template_content:
            return ""
            
        def replace_match(match):
            key = match.group(1).strip().lower()
            return replace_dict.get(key, match.group(0))
            
        return self.pattern.sub(replace_match, self.template_content)


class MailDistributionContext:
    """엑셀 파일과 MSDS 폴더를 읽어 동적 메일 배포 정보를 구성하고 자가 청소 로그를 관리하는 단일 Deep Module."""
    def __init__(self, excel_path: str, msds_dir: Optional[str] = None, default_sender: Optional[str] = None):
        self.excel_path = excel_path
        self.msds_dir = msds_dir
        self.default_sender = default_sender
        self.logger = logging.getLogger("MailDistributionContext")
        self.logger.setLevel(logging.INFO)
        
        # 중복 핸들러 방지
        if not self.logger.handlers:
            self._init_logger()
            
    def _init_logger(self):
        log_dir = "logs"
        if not os.path.exists(log_dir):
            os.makedirs(log_dir)
            
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        log_file = os.path.join(log_dir, f"distribution_{timestamp}.log")
        
        file_handler = logging.FileHandler(log_file, encoding="utf-8")
        file_handler.setLevel(logging.INFO)
        
        stream_handler = logging.StreamHandler()
        stream_handler.setLevel(logging.INFO)
        
        formatter = logging.Formatter("[%(asctime)s] [%(levelname)s] %(message)s", datefmt="%Y-%m-%d %H:%M:%S")
        file_handler.setFormatter(formatter)
        stream_handler.setFormatter(formatter)
        
        self.logger.addHandler(file_handler)
        self.logger.addHandler(stream_handler)
        
        self.logger.info(f"감사 로그가 생성되었습니다: {os.path.abspath(log_file)}")

    def _cleanup_old_logs(self, retention_days: int = 7):
        log_dir = "logs"
        if not os.path.exists(log_dir):
            return
            
        self.logger.info(f"기한이 경과한 감사 로그 자동 정리를 시도합니다 (보존 기간: {retention_days}일)...")
        now = datetime.datetime.now()
        threshold = now - datetime.timedelta(days=retention_days)
        
        try:
            files = os.listdir(log_dir)
            cleaned_count = 0
            for f in files:
                if f.startswith("distribution_") and f.endswith(".log"):
                    fpath = os.path.join(log_dir, f)
                    mtime = datetime.datetime.fromtimestamp(os.path.getmtime(fpath))
                    if mtime < threshold:
                        os.remove(fpath)
                        self.logger.info(f"   -> [삭제 완료] 보존 기한 만료 로그 삭제: {f}")
                        cleaned_count += 1
            if cleaned_count > 0:
                self.logger.info(f"총 {cleaned_count}개의 오래된 감사 로그 파일이 정리되었습니다.")
            else:
                self.logger.info("정리할 대상 오래된 감사 로그 파일이 없습니다.")
        except Exception as e:
            self.logger.warning(f"감사 로그 파일 자동 정리 과정 중 경고 발생: {str(e)}")

    def load_jobs(self) -> List[MailJob]:
        """엑셀 파싱 및 MSDS 매칭, 본문 표 자동 조립을 수행하여 발송 정보 객체 목록을 로드함."""
        self._cleanup_old_logs()
        self.logger.info(f"엑셀 데이터 로딩을 개시합니다: {self.excel_path}")
        
        try:
            parser = ExcelParser(self.excel_path)
        except Exception as e:
            self.logger.error(f"엑셀 로드에 실패하여 프로그램을 중단합니다: {str(e)}")
            raise
            
        recipients = parser.get_recipients()
        table_headers = parser.get_table_headers()
        html_table_renderer = HtmlTableRenderer(table_headers)
        
        jobs: List[MailJob] = []
        
        search_dir = self.msds_dir if self.msds_dir else "."
        all_files = []
        if os.path.exists(search_dir) and os.path.isdir(search_dir):
            try:
                all_files = os.listdir(search_dir)
            except Exception as e:
                self.logger.warning(f"MSDS 폴더 '{search_dir}'를 스캔하는 과정에서 에러 발생: {str(e)}")
        else:
            self.logger.warning(f"지정된 MSDS 폴더가 존재하지 않거나 디렉토리가 아닙니다: {search_dir}")

        self.logger.info(f"총 {len(recipients)}명의 메일 명세 분석을 시작합니다.")
        
        for idx, rec in enumerate(recipients, 1):
            to_addr = rec["to"]
            cc_addr = rec["cc"]
            bcc_addr = rec["bcc"]
            subject = rec["subject"]
            attachments = rec["attachments"][:]
            
            self.logger.info(f"[{idx}/{len(recipients)}] 수신자 '{to_addr}' 분석 중...")
            
            # 1. 취급 제품 테이블 빌드 및 MSDS 접두사 매칭 스캔
            html_table = ""
            table_rows = parser.get_table_data(to_addr)
            if table_rows and table_headers:
                html_table = html_table_renderer.render(table_rows)

            msds_missing = False
            missing_prod_code = ""
            
            # 제품코드 접두사 자동 수집
            for row in table_rows:
                prod_code = None
                for k, v in row.items():
                    if k.strip().lower() == "제품코드":
                        prod_code = str(v).strip()
                        break
                
                if prod_code:
                    matched_in_dir = False
                    # ponytail: 제품코드 간의 모호성을 제거하기 위해 무조건 '제품코드_' 접두사 패턴으로 매칭
                    for filename in all_files:
                        if filename.lower().startswith(prod_code.lower() + "_"):
                            full_path = os.path.abspath(os.path.join(search_dir, filename))
                            if full_path not in attachments:
                                attachments.append(full_path)
                                self.logger.info(f"   -> [자동 첨부] 제품코드 '{prod_code}' 매칭 파일: {filename}")
                            matched_in_dir = True
                    
                    if not matched_in_dir:
                        msds_missing = True
                        missing_prod_code = prod_code

            # 개별 검증 1: MSDS 첨부파일 실종 시 스킵 정책
            if msds_missing:
                self.logger.error(f"   -> [오류/건너뜀] 제품코드 '{missing_prod_code}'에 해당하는 MSDS 서류가 폴더 내에 실종되었습니다. 이 수신자는 스킵합니다.")
                continue

            # 2. 본문 치환 엔진 가동 및 표 병합
            final_body = ""
            if rec["body"]:
                body_html = rec["body"].replace("\n", "<br>")
                replace_dict = rec["replace_dict"]
                replace_dict["table"] = html_table
                
                temp_engine = TemplateEngine(body_html)
                final_body = temp_engine.render(replace_dict)
                
                if "{{table}}" not in body_html and html_table:
                    final_body += "<br><br>" + html_table
            else:
                final_body = html_table if html_table else "본문 내용이 없습니다."

            # 3. 맑은 고딕 11pt 통합 폰트 및 스타일 랩핑
            font_wrapped_body = f'<div style="font-family: \'맑은 고딕\', \'Malgun Gothic\', sans-serif; font-size: 11.0pt; color: #000000; line-height: 1.6;">{final_body}</div>'

            # 개별 검증 2: 예약 일시 유효성 검사 (포맷 검증)
            deferred_time = rec.get("deferred_time")
            if deferred_time:
                # 문자열 타입의 일시면 정규식 검증 시도
                if isinstance(deferred_time, str) and deferred_time.strip():
                    deferred_time = deferred_time.strip()
                    # YYYY-MM-DD HH:MM:SS 규격 검증
                    pattern = r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$"
                    if not re.match(pattern, deferred_time):
                        self.logger.error(f"   -> [오류/건너뜀] 예약 발송 시간 '{deferred_time}'의 날짜 포맷이 올바르지 않습니다 (YYYY-MM-DD HH:MM:SS 규격 필요). 이 수신자는 스킵합니다.")
                        continue
            
            # 발신인 계정 결정
            from_addr = rec.get("from") if rec.get("from") else self.default_sender
            
            # 정상 로드된 메일 작업 추가
            jobs.append(MailJob(
                to_addr=to_addr,
                subject=subject,
                body=font_wrapped_body,
                cc_addr=cc_addr if cc_addr else None,
                bcc_addr=bcc_addr if bcc_addr else None,
                attachments=attachments,
                from_addr=from_addr if from_addr else None,
                deferred_time=deferred_time
            ))
            self.logger.info("   -> [분석 완료] 발송 대기열 추가 성공")

        self.logger.info(f"분석 요약: 전체 {len(recipients)}건 중 {len(jobs)}건 변환 성공 (실패/스킵: {len(recipients) - len(jobs)}건)")
        return jobs

    def close(self):
        """로거 파일 핸들러의 스트림 락을 안전하게 해제합니다."""
        handlers = self.logger.handlers[:]
        for handler in handlers:
            try:
                handler.close()
                self.logger.removeHandler(handler)
            except Exception:
                pass
